from pathlib import Path
from shutil import copyfile
from copy import deepcopy
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SOURCE = Path(r"C:\Users\Faiz6\Downloads\template_5.docx")
OUTPUT = Path(r"C:\Users\Faiz6\OneDrive\Desktop\AI-Interview-Preparation-Platform\docs\Submission-5-Implementation-and-Testing.docx")
FONT = "Times New Roman"


def set_run(run, size=12, bold=False, italic=False, color="000000", font=FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])


def set_table_borders(table, color="808080", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)
        borders.append(element)


def clear_body(doc):
    body = doc._element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.0
    set_run(p.add_run(text), size=14, bold=True)


def add_heading(doc, text, level=1, page_break=False):
    p = doc.add_paragraph()
    if page_break:
        p.paragraph_format.page_break_before = True
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    set_run(p.add_run(text), size=14 if level <= 2 else 12, bold=True)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), size=12, bold=True)
        set_run(p.add_run(text[len(bold_lead):]), size=12)
    else:
        set_run(p.add_run(text), size=12)
    return p


def add_bullets(doc, items, size=12, space_after=4):
    for item in items:
        p = doc.add_paragraph(style="List Paragraph")
        if BULLET_NUM_PR is not None:
            p._p.get_or_add_pPr().append(deepcopy(BULLET_NUM_PR))
        p.paragraph_format.left_indent = Inches(0.45)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.05 if size < 12 else 1.1
        set_run(p.add_run(item), size=size)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for index, heading in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, "D9E2F3")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(heading), size=10, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            set_run(p.add_run(str(value)), size=10)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_code(doc, title, code):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(title), size=11, bold=True)
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="BFBFBF", size="4")
    set_repeat_header(table.rows[0])
    shade_cell(table.cell(0, 0), "F2F2F2")
    set_table_geometry(table, [9000])
    code_p = table.cell(0, 0).paragraphs[0]
    code_p.paragraph_format.space_after = Pt(0)
    code_p.paragraph_format.line_spacing = 1.0
    set_run(code_p.add_run(code.strip()), size=8.5, font="Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


copyfile(SOURCE, OUTPUT)
doc = Document(OUTPUT)
BULLET_NUM_PR = deepcopy(doc.paragraphs[3]._p.pPr.numPr) if doc.paragraphs[3]._p.pPr.numPr is not None else None
clear_body(doc)
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
normal.font.size = Pt(12)

add_title(doc, "5. IMPLEMENTATION AND TESTING")

add_heading(doc, "5.1 Implementation Approach")
add_heading(doc, "5.1.1 Introduction", level=2)
add_body(doc, "Module 5 implements the interactive answer and evaluation stage of the AI Interview Preparation Platform. After Module 4 creates role-specific questions, the candidate can type an answer or dictate it using the microphone. Each response is validated, associated with the authenticated user, stored in MongoDB Atlas, evaluated through Gemini AI or a permanent local fallback, and presented with scores and improvement feedback.")
add_body(doc, "The implementation follows the existing MVC architecture. React manages the interview interface; Axios sends authenticated requests; Express routes requests to controllers; service modules perform AI or rule-based evaluation; Mongoose validates and stores answers; and the final report aggregates all question-level evaluations.")
add_bullets(doc, [
    "Typed and voice-to-text answer input with editable transcripts.",
    "User, interview and question ownership validation through JWT-protected APIs.",
    "Five evaluation scores: correctness, confidence, communication, technical accuracy and overall performance.",
    "Strengths, weaknesses and suggestions for every submitted answer.",
    "Gemini evaluation with a deterministic local fallback when the external service is unavailable.",
    "Interview completion validation and a persistent final performance report.",
], size=11, space_after=2)

add_heading(doc, "5.1.2 Input and Output Design Implementation", level=2)
add_body(doc, "The main input is the candidate's response to the displayed interview question. The answer may be entered directly in a textarea or generated as text through browser speech recognition. The interface limits each response to 5,000 characters, displays a live counter, prevents blank submissions and disables navigation while voice recognition or saving is in progress.")
add_table(doc, ["Input / Output", "Implementation", "Validation or Purpose"], [
    ("Typed answer", "React controlled textarea", "2-5,000 characters; editable before submission"),
    ("Voice answer", "Web Speech API transcript", "Microphone permission, browser support and recognition errors handled"),
    ("Answer request", "questionId and answer text", "JWT identifies the user; URL identifies the interview"),
    ("Evaluation output", "Scores and feedback lists", "Each score is constrained between 0 and 100"),
    ("Final report", "Interview averages and detailed review", "Available only after every question is evaluated"),
], [1700, 3300, 4000])

add_heading(doc, "5.1.3 Database Implementation", level=2)
add_body(doc, "MongoDB Atlas provides permanent, user-specific storage. The Answer model references the User, Interview and Question collections. A compound unique index on userId, interviewId and questionId ensures that one candidate has only one current response per interview question. Re-submission updates the same document and resets the previous evaluation before calculating new feedback.")
add_body(doc, "The Interview model stores the final status, overall score, completion time and a structured performanceReport object. This report contains category averages and combined strengths, weaknesses and suggestions, allowing Module 6 to generate history and analytics without recalculating completed interviews.")

add_heading(doc, "5.1.4 Collection Structures", level=2)
add_table(doc, ["Collection", "Important fields", "Relationship / Role"], [
    ("Users", "_id, name, email, password", "Owns resumes, interviews and answers"),
    ("Resumes", "userId, analysis, interviewPreferences", "Supplies candidate context and target role"),
    ("Interviews", "userId, resumeId, role, score, performanceReport", "Stores configuration and final result"),
    ("Questions", "interviewId, order, question, expectedPoints", "Provides the evaluation rubric"),
    ("Answers", "userId, interviewId, questionId, answer, scores, feedback", "Stores each evaluated response"),
], [1500, 4100, 3400])

add_heading(doc, "5.1.5 Code Modules", level=2)
add_table(doc, ["Layer", "File", "Responsibility"], [
    ("Model", "server/models/Answer.js", "Answer constraints, evaluation fields and unique index"),
    ("Controller", "server/controllers/answerController.js", "Save, evaluate, retrieve and complete interviews"),
    ("Route", "server/routes/answerRoutes.js", "Protected answer and completion endpoints"),
    ("AI service", "server/services/geminiAnswerEvaluator.js", "Structured contextual evaluation through Gemini"),
    ("Fallback", "server/services/localAnswerEvaluator.js", "Rule-based evaluation available without Gemini"),
    ("Orchestration", "server/services/answerEvaluationService.js", "Gemini attempts and automatic fallback selection"),
    ("Frontend", "client/src/pages/MockInterview.jsx", "Text input, microphone input, saving and feedback display"),
    ("Frontend", "client/src/pages/InterviewReport.jsx", "Final category scores and detailed answer review"),
], [1500, 3300, 4200])

add_heading(doc, "5.1.6 System Implementation", level=2)
add_body(doc, "When the candidate selects Save and Continue, the frontend sends POST /api/answers/:interviewId. The controller validates MongoDB identifiers, checks that the interview belongs to the authenticated user, verifies that the question belongs to that interview and saves the normalized answer. It then invokes the evaluation service and stores the result before returning it to the interface.")
add_body(doc, "On the final question, POST /api/answers/:interviewId/complete verifies that the number of completed evaluations equals the number of interview questions. It calculates category averages, creates the combined report, changes the status to completed and navigates the candidate to /interview/:interviewId/report. Opening this route after a refresh reloads the persisted interview and answers.")

add_heading(doc, "5.1.7 Project Summary", level=2)
add_body(doc, "Module 5 converts the question-display prototype into a complete mock-interview experience. It supports accessible answer entry, secure persistence, resilient evaluation and actionable reporting. The implementation remains modular so the AI provider, scoring rules or input method can be changed without rewriting the entire application.")

add_heading(doc, "5.2 Coding Details and Code Efficiency")
add_body(doc, "Representative code from the implemented module is shown below. Only focused excerpts are included so that the report explains the important logic without duplicating entire source files.")

add_code(doc, "A. Answer ownership and duplicate prevention - Answer.js", """
answerSchema.index(
  { userId: 1, interviewId: 1, questionId: 1 },
  { unique: true }
);

answer: {
  type: String,
  required: true,
  trim: true,
  minlength: 2,
  maxlength: 5000,
}
""")
add_body(doc, "The compound index supports fast ownership queries and prevents duplicate responses. Mongoose validation rejects invalid text before it reaches permanent storage.")

add_code(doc, "B. Browser voice-to-text initialization - MockInterview.jsx", """
const SpeechRecognition =
  window.SpeechRecognition || window.webkitSpeechRecognition;

const recognition = new SpeechRecognition();
recognition.lang = navigator.language || "en-US";
recognition.continuous = true;
recognition.interimResults = true;
recognition.start();
""")
add_body(doc, "The browser converts speech into editable text. The platform stores only the transcript, not an audio file, reducing database storage and privacy exposure. Unsupported-browser, permission, microphone and recognition-service errors are displayed safely.")

add_code(doc, "C. Gemini with permanent local fallback - answerEvaluationService.js", """
for (const model of GEMINI_MODELS) {
  try {
    const evaluation = await evaluateWithGemini({ ...options, model });
    return { ...evaluation, source: "gemini", model };
  } catch (error) {
    lastError = error;
  }
}

return { ...evaluateLocally(options), source: "local", model: null };
""")
add_body(doc, "The orchestration service isolates external API failure from the controller. Gemini results are schema-validated; if all attempts fail, the local evaluator returns an identical response structure so the UI continues to work.")

add_code(doc, "D. Weighted local overall score - localAnswerEvaluator.js", """
const overall = clamp(
  correctness * 0.35 +
  confidence * 0.20 +
  communication * 0.25 +
  technicalAccuracy * 0.20
);
""")
add_body(doc, "Correctness receives the largest weight, while communication, confidence and role-specific accuracy also contribute. The clamp function guarantees an integer score between 0 and 100.")

add_code(doc, "E. Final report aggregation - answerController.js", """
const average = (field) => Math.round(
  answers.reduce((total, item) => total + item.scores[field], 0)
  / answers.length
);

interview.status = "completed";
interview.score = average("overall");
interview.completedAt = interview.completedAt || new Date();
await interview.save();
""")

add_heading(doc, "5.2.1 Code Efficiency Measures", level=2)
add_bullets(doc, [
    "MVC separation keeps database, request-handling, evaluation and interface responsibilities independent.",
    "Database indexes improve common lookups and enforce integrity without repeated application-side scans.",
    "findOneAndUpdate with upsert updates an existing answer or creates it in one database operation.",
    "Expected answer points stored with each question avoid regenerating an evaluation rubric.",
    "Structured Gemini JSON validation prevents malformed AI data from entering MongoDB.",
    "Bounded feedback arrays and answer length limits prevent uncontrolled document growth.",
    "The local evaluator uses sets for unique tokens and direct keyword lookup, keeping processing proportional to answer length.",
    "React stores answers by question ID, enabling constant-time retrieval when navigating between questions.",
], size=11, space_after=2)

add_heading(doc, "5.3 Testing Approach")
add_body(doc, "Testing was performed at unit, integration and complete-system levels. Automated syntax checks, module import checks, ESLint and Vite production builds were combined with browser tests for answer entry, voice transcription, persistence, evaluation and report navigation. Test data used a five-question interview to keep repeated end-to-end checks efficient.")

add_heading(doc, "5.3.1 Unit Testing", level=2)
doc.add_paragraph().paragraph_format.space_after = Pt(1)
add_body(doc, "Unit-level checks focused on isolated validation and scoring logic. Node's syntax checker verified every new backend module. The local evaluator was executed with a controlled technical question and answer, and the test confirmed that it returned five integer scores, strengths, weaknesses and suggestions within the required bounds.")
add_table(doc, ["Unit", "Test performed", "Expected result", "Result"], [
    ("Answer validation", "Submit fewer than 2 or more than 5,000 characters", "HTTP 400 with clear message", "Pass"),
    ("ObjectId validation", "Invalid interview or question ID", "Request rejected before query", "Pass"),
    ("Local scoring", "Evaluate controlled answer text", "Scores are integers from 0 to 100", "Pass"),
    ("Gemini validation", "Malformed or incomplete response", "Reject response and use fallback", "Pass"),
    ("Frontend code", "Run ESLint", "No lint errors", "Pass"),
], [1700, 3100, 3000, 1200])

add_heading(doc, "5.3.2 Integration Testing", level=2)
doc.add_paragraph().paragraph_format.space_after = Pt(1)
add_body(doc, "Integration testing verified communication between React, Axios, JWT middleware, controllers, services and MongoDB Atlas. A saved answer was revisited and remained visible after navigation and refresh. Changing and resubmitting it updated the same database record. The evaluation panel displayed the source and persisted feedback returned by the protected API.")
add_table(doc, ["Integration", "Components", "Observed result"], [
    ("Authentication", "Axios + JWT middleware + User model", "Only authenticated users access answer APIs"),
    ("Answer persistence", "MockInterview + controller + Answer model", "Answer survives navigation and refresh"),
    ("Evaluation", "Controller + Gemini/local services + MongoDB", "Scores and feedback saved with source"),
    ("Completion", "Questions + Answers + Interview model", "Incomplete interview rejected; complete report stored"),
    ("Report reload", "React Router + interview/answer APIs", "Completed report opens after refresh"),
], [1900, 3500, 3600])

add_heading(doc, "5.3.3 System Testing", level=2)
doc.add_paragraph().paragraph_format.space_after = Pt(1)
add_body(doc, "System testing followed the complete user workflow: register, log in, upload a resume, generate an interview, answer using text and microphone, inspect individual feedback, complete all questions and open the final report. The backend remained connected to MongoDB Atlas, and the application displayed friendly empty, loading and error states rather than exposing technical database details.")
add_table(doc, ["System test", "Acceptance criterion", "Result"], [
    ("Typed-answer workflow", "Save, navigate back and reload the stored response", "Pass"),
    ("Voice-to-text workflow", "Start, stop, edit and submit transcript in Chrome", "Pass"),
    ("Fallback availability", "Local questions/evaluation returned when Gemini unavailable", "Pass"),
    ("Final performance report", "All category averages and detailed review displayed", "Pass"),
    ("Production frontend", "Vite production build completes", "Pass"),
    ("Backend structure", "Syntax and import checks complete", "Pass"),
], [2500, 5000, 1500])

add_heading(doc, "5.4 Modifications and Improvements")
add_heading(doc, "5.4.1 Modifications Completed", level=2)
add_bullets(doc, [
    "Replaced the original practice-note placeholder with a real controlled answer field.",
    "Changed the technical user message from 'Saved in MongoDB' to the clearer 'Answer saved'.",
    "Added browser voice-to-text with start/stop controls, interim transcripts and permission errors.",
    "Added ownership checks so users cannot read or submit answers for another user's interview.",
    "Added answer upsert behavior to prevent duplicate documents when an answer is edited.",
    "Added Gemini evaluation and a permanent local fallback with a visible source label.",
    "Added final-interview validation, aggregate scoring and a refresh-safe performance-report route.",
    "Added responsive score cards and feedback panels for desktop and mobile screens.",
], size=11, space_after=2)

add_heading(doc, "5.4.2 Current Limitations", level=2)
add_bullets(doc, [
    "Web Speech API support and recognition quality vary by browser, language and external speech service availability.",
    "Only the transcript is analysed; actual voice tone, pauses, facial expression and body language are not measured.",
    "The local evaluator is deterministic and less context-aware than Gemini.",
    "Gemini availability remains subject to network access, model availability and API quota.",
    "Keyword-based resume parsing may miss unusual resume layouts or uncommon terminology.",
], size=11, space_after=2)

add_heading(doc, "5.4.3 Future Improvements", level=2)
add_bullets(doc, [
    "Add server-side multilingual speech transcription for consistent browser support.",
    "Add optional audio/video recording with explicit consent and secure retention controls.",
    "Measure delivery pace, filler words and pauses when audio analysis is available.",
    "Add coding questions with an editor and automated test cases.",
    "Use Module 6 history to generate personalized practice plans from repeated weaknesses.",
    "Add automated backend API tests with a dedicated test database and continuous integration.",
], size=11, space_after=2)

add_heading(doc, "5.4.4 Conclusion", level=2)
add_body(doc, "Module 5 successfully implements the core interview-answering and evaluation workflow. Its secure MongoDB relationships, hybrid evaluation strategy and persisted performance report provide a reliable foundation for interview history and analytics. The module meets the functional requirement of returning scores, strengths, weaknesses and practical suggestions for every answer while remaining operational during external AI failure.")

doc.core_properties.title = "Submission 5 - Implementation and Testing"
doc.core_properties.subject = "AI Interview Preparation Platform - Module 5"
doc.core_properties.author = "Student"
doc.core_properties.keywords = "AI interview, implementation, testing, MongoDB, React, Express, Gemini"
doc.save(OUTPUT)
print(OUTPUT)
