from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Faiz6\OneDrive\Desktop\AI-Interview-Preparation-Platform")
OUTPUT = ROOT / "docs" / "AI-Interview-Platform-Data-Design-Report.docx"
SCREENSHOTS = {
    "login": Path(r"C:\Users\Faiz6\AppData\Local\Temp\codex-clipboard-1d8fbcc8-4196-4bba-bf67-a975d3b83d70.png"),
    "dashboard": Path(r"C:\Users\Faiz6\AppData\Local\Temp\codex-clipboard-a505ec23-7d86-40c2-ba3d-3f09c398d8eb.png"),
    "setup": Path(r"C:\Users\Faiz6\AppData\Local\Temp\codex-clipboard-76c8b0c2-9728-4b72-8f0c-062eb146c660.png"),
    "interview": Path(r"C:\Users\Faiz6\AppData\Local\Temp\codex-clipboard-ebce8d70-1f73-41b4-b5fd-76aea0e4a478.png"),
}

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
BORDER = "CBD5E1"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_run_font(run, size=11, color="000000", bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (fld_char, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=9, color=MUTED)


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, bold_label=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.10
    if bold_label and text.startswith(bold_label):
        label = paragraph.add_run(bold_label)
        set_run_font(label, bold=True)
        run = paragraph.add_run(text[len(bold_label):])
        set_run_font(run)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing = 1.167
        run = paragraph.add_run(item)
        set_run_font(run)


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, heading in enumerate(headers):
        set_cell_shading(header.cells[index], LIGHT_GRAY)
        paragraph = header.cells[index].paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(heading)
        set_run_font(run, size=10, color=NAVY, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            set_run_font(run, size=9.5)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    paragraph = cell.paragraphs[0]
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, color=DARK_BLUE, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_screenshot(doc, path, caption, width=6.25, crop_bottom=0):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("title", caption)
    shape._inline.docPr.set("descr", caption)
    if crop_bottom:
        src_rect = OxmlElement("a:srcRect")
        src_rect.set("b", str(int(crop_bottom * 100000)))
        shape._inline.graphic.graphicData.pic.blipFill.insert(1, src_rect)
        shape.height = int(shape.height * (1 - crop_bottom))
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(4)
    caption_p.paragraph_format.space_after = Pt(12)
    caption_p.paragraph_format.keep_with_next = True
    caption_run = caption_p.add_run(caption)
    set_run_font(caption_run, size=9, color=MUTED, italic=True)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_header_footer(section):
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run("AI Interview Preparation Platform | Field Project 2026-27")
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Data Design Report  |  Page ")
    set_run_font(run, size=9, color=MUTED)
    add_field(paragraph, "PAGE")


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("TYSD FIELD PROJECT | ACADEMIC YEAR 2026-27")
    set_run_font(run, size=11, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("AI Interview Preparation Platform")
    set_run_font(run, size=28, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("Basic Modules, Data Design, Schema Design,\nand Data Integrity & Constraints")
    set_run_font(run, size=15, color=DARK_BLUE)

    metadata = [
        ("Student Name", "________________________________"),
        ("Roll Number", "________________________________"),
        ("Class / Division", "________________________________"),
        ("Project Guide", "________________________________"),
        ("College", "________________________________"),
        ("Submission Date", "________________________________"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in metadata:
        cells = table.add_row().cells
        cells[0].paragraphs[0].add_run(label)
        cells[1].paragraphs[0].add_run(value)
        set_cell_shading(cells[0], LIGHT_GRAY)
        for run in cells[0].paragraphs[0].runs:
            set_run_font(run, size=10, color=NAVY, bold=True)
        for run in cells[1].paragraphs[0].runs:
            set_run_font(run, size=10)
    set_repeat_table_header(table.rows[0])
    set_table_geometry(table, [2700, 6660])

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(26)
    run = note.add_run("Submission: 4.1 Basic Modules | 4.2 Data Design | 4.2.1 Schema Design | 4.2.2 Data Integrity and Constraints")
    set_run_font(run, size=9.5, color=MUTED, italic=True)
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        configure_header_footer(section)

    properties = doc.core_properties
    properties.title = "AI Interview Preparation Platform - Data Design Report"
    properties.subject = "Field Project submission for basic modules, schema design, and data integrity"
    properties.author = "Student"
    properties.keywords = "AI interview, MongoDB, schema design, data integrity, field project"

    add_cover(doc)

    add_heading(doc, "Executive Summary", 1)
    add_body(doc, "The AI Interview Preparation Platform is a full-stack web application that helps users prepare for job interviews through authenticated accounts, resume analysis, role-based question generation, and guided mock interviews. React and Vite provide the user interface, while Node.js and Express implement the API. MongoDB Atlas is the permanent data store, and Mongoose defines collection schemas and validation rules.")
    add_callout(doc, "Database purpose", "MongoDB stores permanent user, resume, interview, and question records. Browser localStorage is used only as a temporary frontend cache and is not the authoritative data source.")

    add_heading(doc, "Document Scope", 2)
    add_bullets(doc, [
        "4.1 Basic Modules - functional components currently implemented in the platform.",
        "4.2 Data Design - database structure, collection ownership, and data flow.",
        "4.2.1 Schema Design - field definitions for Users, Resumes, Interviews, and Questions.",
        "4.2.2 Data Integrity and Constraints - validation, security, uniqueness, and ownership controls.",
    ])

    add_heading(doc, "4.1 Basic Modules", 1)
    module_rows = [
        ("1", "Registration and Login", "Creates accounts, hashes passwords, verifies credentials, and issues JWTs.", "Implemented"),
        ("2", "Resume Upload", "Accepts authenticated PDF/DOCX uploads up to 5 MB using Multer.", "Implemented"),
        ("3", "Resume Analysis", "Extracts text, skills, education, experience, and keywords and stores the result.", "Implemented"),
        ("4", "Resume Profile Editing", "Allows users to correct target role, experience level, skills, education, experience, and keywords.", "Implemented"),
        ("5", "Question Generation", "Uses Gemini first and a permanent role-aware local fallback if the API is unavailable.", "Implemented"),
        ("6", "Mock Interview", "Displays saved questions one at a time with type, difficulty, source, and progress.", "Implemented"),
        ("7", "Answer Evaluation", "Will evaluate correctness, confidence, communication, and technical accuracy.", "Planned"),
        ("8", "Analytics and History", "Will show previous interviews, scores, progress charts, and suggestions.", "Planned"),
    ]
    add_table(doc, ["No.", "Module", "Responsibility", "Status"], module_rows, [600, 1900, 5260, 1600])

    add_heading(doc, "Module Interaction", 2)
    add_body(doc, "The modules execute in a controlled sequence: registration -> login -> resume upload -> resume analysis -> profile review/editing -> interview configuration -> question generation -> mock interview. Evaluation and analytics will extend this sequence in later modules.")

    add_heading(doc, "4.2 Data Design", 1)
    add_body(doc, "MongoDB Atlas was selected because the project stores nested resume-analysis data, evolving interview records, and user-specific documents. Mongoose provides application-level schemas, validation, indexes, references, defaults, and data-type enforcement.")

    add_heading(doc, "Collection Relationship Model", 2)
    relationship_rows = [
        ("Users", "Resumes", "One-to-many", "Resumes.userId -> Users._id"),
        ("Users", "Interviews", "One-to-many", "Interviews.userId -> Users._id"),
        ("Resumes", "Interviews", "One-to-many", "Interviews.resumeId -> Resumes._id"),
        ("Interviews", "Questions", "One-to-many", "Questions.interviewId -> Interviews._id"),
    ]
    add_table(doc, ["Parent", "Child", "Cardinality", "Reference"], relationship_rows, [1800, 1800, 1800, 3960])
    add_callout(doc, "Ownership rule", "Every protected read or update filters by the authenticated user ID. A user cannot retrieve or edit another user's resume or interview by changing a URL identifier.")

    add_heading(doc, "4.2.1 Schema Design", 1)
    add_heading(doc, "Users Collection", 2)
    add_body(doc, "Stores registered user identities. Password values contain bcrypt hashes and are excluded from normal model queries.")
    users_rows = [
        ("_id", "ObjectId", "Yes", "MongoDB document identifier"),
        ("name", "String", "Yes", "Trimmed; 2-80 characters"),
        ("email", "String", "Yes", "Unique, lowercase, and trimmed"),
        ("password", "String", "Yes", "bcrypt hash; select:false"),
        ("createdAt", "Date", "Yes", "Defaults to current date/time"),
    ]
    add_table(doc, ["Field", "Type", "Required", "Description / Constraint"], users_rows, [1700, 1400, 1200, 5060])

    add_heading(doc, "Resumes Collection", 2)
    add_body(doc, "Stores file metadata, the authenticated owner, interview preferences, extracted text, and editable structured analysis.")
    resume_rows = [
        ("_id", "ObjectId", "Yes", "Resume identifier"),
        ("userId", "ObjectId ref User", "Owner", "Connects the resume to its user"),
        ("fileName", "String", "Yes", "Original uploaded filename"),
        ("storedFileName", "String", "Yes", "Unique safe server filename"),
        ("filePath", "String", "Yes", "Relative path under server/uploads"),
        ("mimeType", "String", "Yes", "PDF or DOCX MIME type"),
        ("interviewPreferences.role", "String", "Defaulted", "Editable target job role"),
        ("interviewPreferences.experienceLevel", "String", "Defaulted", "Fresher, 1-2, 3-5, or 5+ Years"),
        ("interviewPreferences.difficulty", "String", "Defaulted", "Beginner, Intermediate, or Advanced"),
        ("analysis.skills", "String[]", "Default []", "Recognized and user-edited skills"),
        ("analysis.education", "String[]", "Default []", "Education lines"),
        ("analysis.experience", "String[]", "Default []", "Experience lines"),
        ("analysis.keywords", "String[]", "Default []", "High-value resume terms"),
        ("analysis.rawText", "String", "Yes", "Complete extracted document text"),
        ("uploadedAt", "Date", "Defaulted", "Upload timestamp"),
    ]
    add_table(doc, ["Field", "Type", "Required", "Description / Constraint"], resume_rows, [2450, 1700, 1250, 3960])

    add_heading(doc, "Interviews Collection", 2)
    interview_rows = [
        ("_id", "ObjectId", "Yes", "Interview identifier"),
        ("userId", "ObjectId ref User", "Yes", "Indexed owner reference"),
        ("resumeId", "ObjectId ref Resume", "Yes", "Source resume reference"),
        ("role", "String", "Yes", "Trimmed; maximum 120 characters"),
        ("type", "String enum", "Yes", "HR, Technical, Behavioral, or Mixed"),
        ("difficulty", "String enum", "Yes", "Beginner, Intermediate, or Advanced"),
        ("questionCount", "Number", "Yes", "Minimum 5; maximum 15"),
        ("generationSource", "String enum", "Default local", "gemini or local"),
        ("generationModel", "String/null", "No", "Successful Gemini model identifier"),
        ("status", "String enum", "Default", "not_started, in_progress, or completed"),
        ("score", "Number/null", "No", "Future score constrained to 0-100"),
        ("feedback", "String", "Default empty", "Future evaluation feedback"),
        ("createdAt", "Date", "Defaulted", "Creation timestamp"),
    ]
    add_table(doc, ["Field", "Type", "Required", "Description / Constraint"], interview_rows, [2250, 1700, 1300, 4110])

    add_heading(doc, "Questions Collection", 2)
    question_rows = [
        ("_id", "ObjectId", "Yes", "Question identifier"),
        ("interviewId", "ObjectId ref Interview", "Yes", "Indexed parent interview"),
        ("order", "Number", "Yes", "Minimum 1; unique within interview"),
        ("question", "String", "Yes", "Trimmed question text"),
        ("type", "String enum", "Yes", "HR, Technical, or Behavioral"),
        ("difficulty", "String enum", "Yes", "Beginner, Intermediate, or Advanced"),
        ("skill", "String/null", "No", "Related skill or competency"),
        ("expectedPoints", "String[]", "Default []", "Reference points for future evaluation"),
    ]
    add_table(doc, ["Field", "Type", "Required", "Description / Constraint"], question_rows, [2250, 1800, 1300, 4010])

    add_heading(doc, "4.2.2 Data Integrity and Constraints", 1)
    integrity_rows = [
        ("Identity", "Unique email index", "Prevents duplicate registered accounts."),
        ("Credential security", "bcrypt hashing", "Plain-text passwords are never stored."),
        ("Authorization", "JWT + ownership filter", "Protects private routes and prevents cross-user access."),
        ("File integrity", "Extension, MIME, and size validation", "Only PDF/DOCX files up to 5 MB are accepted."),
        ("Resume editing", "Array cleaning and item limits", "Removes blanks/duplicates and limits field sizes."),
        ("Interview type", "Mongoose enum", "Rejects types outside HR, Technical, Behavioral, and Mixed."),
        ("Difficulty", "Mongoose enum", "Rejects levels outside Beginner, Intermediate, and Advanced."),
        ("Question count", "min 5 / max 15", "Prevents invalid or excessive generation requests."),
        ("Question order", "Compound unique index", "One order number per question within an interview."),
        ("Score", "min 0 / max 100", "Prevents invalid future evaluation scores."),
        ("AI response", "Strict JSON schema + validation", "Rejects malformed, duplicate, or wrong-type Gemini output."),
        ("Availability", "Model failover + local fallback", "External API failures do not crash question generation."),
    ]
    add_table(doc, ["Area", "Constraint / Control", "Integrity Benefit"], integrity_rows, [1900, 3000, 4460])

    add_heading(doc, "Validation Layers", 2)
    add_bullets(doc, [
        "Frontend validation gives immediate feedback for required fields, file type, file size, and form ranges.",
        "Express controllers validate request values and authenticated ownership before database operations.",
        "Mongoose schemas enforce types, required values, defaults, enum lists, numeric ranges, and indexes.",
        "MongoDB Atlas provides persistent storage independent of browser refreshes or application restarts.",
    ])

    add_heading(doc, "End-to-End Data Flow", 1)
    flow_rows = [
        ("1", "Register / Login", "Users", "Create user; return JWT after verification"),
        ("2", "Upload Resume", "Resumes", "Store metadata and parsed analysis"),
        ("3", "Edit Profile", "Resumes", "Save corrected role and structured resume fields"),
        ("4", "Generate Interview", "Interviews", "Store configuration, source, model, and status"),
        ("5", "Generate Questions", "Questions", "Store ordered role-specific questions"),
        ("6", "Mock Interview", "Interviews + Questions", "Load only records owned by current user"),
    ]
    add_table(doc, ["Step", "Application Action", "Collection", "Database Operation"], flow_rows, [700, 2200, 1900, 4560])

    add_heading(doc, "Working Module Evidence", 1)
    add_body(doc, "The following screenshots demonstrate the implemented frontend workflow. Screenshots are included as implementation evidence; personal login details, database credentials, and API keys are not displayed.")
    add_screenshot(doc, SCREENSHOTS["dashboard"], "Figure 1. Resume Analysis Dashboard - extracted skills, education, experience, and keywords", crop_bottom=0.15)
    add_screenshot(doc, SCREENSHOTS["interview"], "Figure 2. Mock Interview - question progress and visible generation source")

    doc.add_page_break()
    add_heading(doc, "Testing Summary", 1)
    test_rows = [
        ("Authentication", "Unknown accounts rejected; registered user receives JWT", "Pass"),
        ("Resume upload", "DOCX upload parsed and saved in MongoDB", "Pass"),
        ("Ownership", "Latest resume/interview filtered by authenticated user", "Pass"),
        ("Question uniqueness", "15 unique questions across types and difficulties", "Pass"),
        ("Difficulty variation", "Distinct Beginner, Intermediate, and Advanced banks", "Pass"),
        ("Gemini integration", "Flash-Lite generated five structured questions", "Pass"),
        ("Fallback", "Local role-aware questions returned when Gemini unavailable", "Pass"),
        ("Frontend quality", "ESLint and Vite production build completed", "Pass"),
    ]
    add_table(doc, ["Test Area", "Evidence", "Result"], test_rows, [2100, 6060, 1200])

    add_heading(doc, "Conclusion", 1)
    add_body(doc, "The current platform implements a modular and secure foundation for AI-assisted interview preparation. MongoDB Atlas provides durable, user-specific storage; Mongoose schemas document and enforce the data model; JWT authentication protects ownership; and the hybrid Gemini/local generator keeps the mock interview usable during external API failures. The schema is prepared for answer evaluation, performance reports, and interview history in subsequent modules.")
    add_callout(doc, "Submission note", "Before submission, replace the cover-page blanks with your student details. If your instructor requires MongoDB Atlas screenshots, add a redacted screenshot of each collection without exposing connection strings, passwords, JWTs, or API keys.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
