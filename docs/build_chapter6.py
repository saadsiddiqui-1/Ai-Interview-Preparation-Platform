from pathlib import Path
from shutil import copyfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SOURCE = Path(r"C:\Users\Faiz6\Downloads\Chapter6.docx")
OUTPUT = Path(r"C:\Users\Faiz6\OneDrive\Desktop\AI-Interview-Preparation-Platform\docs\Chapter-6-Results-and-Discussion.docx")

def font(run, size=12, bold=False, color="000000", name="Times New Roman"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def border(table):
    tbl_pr = table._tbl.tblPr
    edges = tbl_pr.find(qn("w:tblBorders"))
    if edges is None:
        edges = OxmlElement("w:tblBorders")
        tbl_pr.append(edges)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "808080")
        edges.append(node)

def header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)

def shade(cell, value="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), value)
    tc_pr.append(node)

def geometry(table, widths):
    table.autofit = False
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW"); tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[i])); tc_w.set(qn("w:type"), "dxa")

def clear(doc):
    body = doc._element.body
    for item in list(body):
        if item.tag != qn("w:sectPr"): body.remove(item)

def title(doc, value):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    font(p.add_run(value), 14, True)

def heading(doc, value, level=1):
    p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8); p.paragraph_format.space_after = Pt(6)
    font(p.add_run(value), 14 if level == 1 else 12, True)

def body(doc, value):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(.35); p.paragraph_format.space_after = Pt(7); p.paragraph_format.line_spacing = 1.15
    font(p.add_run(value), 12)

def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.45); p.paragraph_format.first_line_indent = Inches(-.2); p.paragraph_format.space_after = Pt(4)
        font(p.add_run("• "), 12, True); font(p.add_run(item), 12)

def table(doc, heads, rows, widths):
    t = doc.add_table(rows=1, cols=len(heads)); border(t); header_row(t.rows[0])
    for i, h in enumerate(heads):
        c=t.rows[0].cells[i]; shade(c); p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); font(p.add_run(h),10,True)
    for data in rows:
        cells=t.add_row().cells
        for i, val in enumerate(data):
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; font(p.add_run(val),10)
    geometry(t,widths); doc.add_paragraph().paragraph_format.space_after=Pt(2)

def screenshot(doc, path, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(path), width=Inches(6.2))
    shape._inline.docPr.set("title", caption)
    shape._inline.docPr.set("descr", caption)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    font(p.add_run(caption), 10, False, "404040")

copyfile(SOURCE, OUTPUT)
doc=Document(OUTPUT); clear(doc)
sec=doc.sections[0]; sec.page_width=Inches(8.27); sec.page_height=Inches(11.69); sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
title(doc,"6. RESULTS AND DISCUSSION")

heading(doc,"6.1 Test Reports")
body(doc,"Module 6 was tested after completing the analytics and interview-history features. The objective was to confirm that a logged-in user can view only their completed interviews, inspect saved performance reports, filter history records and see graph-ready analytics generated from MongoDB data. The tests were performed using the React frontend, Express API, JWT authentication middleware and MongoDB Atlas.")
table(doc,["Test Condition","Input / Action","Expected Result","Actual Result"],[
 ("History API authorization","Call GET /api/interviews/history with JWT","Only logged-in user's completed interviews are returned","Pass - success:true response received"),
 ("Analytics API","Call GET /api/interviews/analytics with JWT","Summary, category averages, trend and suggestions returned","Pass - safe empty response returned when no interview exists"),
 ("Empty analytics","New user with zero completed interviews","Dashboard shows friendly empty state; no crash","Pass"),
 ("Completed interview","Finish every question and save final report","Interview appears in history with score and date","Pass"),
 ("History filter","Search role; choose type and difficulty","Only matching completed interviews displayed","Pass"),
 ("Previous report","Select View from history row","Stored report opens at /interview/:id/report","Pass"),
 ("Pagination","More than ten completed interviews","Previous/Next controls load correct page","Pass"),
 ("Responsive interface","Open analytics on smaller screen","Cards, filters and history remain readable","Pass"),
], [1900, 3000, 3000, 1100])

heading(doc,"6.1.1 Results Discussion",2)
body(doc,"The protected API checks demonstrated that analytics are user-specific. The returned history object contains completed interviews only, while the analytics object contains total interviews, average score, highest score, latest score, improvement value, category averages, score-trend points and repeated suggestions. When the account has no completed interviews, null scores and empty arrays are intentionally returned; this prevents invalid calculations and gives a clean first-use experience.")
body(doc,"After interviews are completed, the dashboard uses the stored performanceReport from Module 5 instead of recalculating every answer. This makes the history screen faster and supports direct navigation to previous reports. The score-trend component uses SVG and does not add an external chart dependency, keeping the frontend lightweight.")

heading(doc,"6.1.2 Module 6 Components",2)
table(doc,["Component","Purpose","Result"],[
 ("GET /api/interviews/history","Returns paginated, filterable history","Protected and user-specific"),
 ("GET /api/interviews/analytics","Calculates overall progress data","Graph-ready JSON response"),
 ("AnalyticsDashboard.jsx","Loads analytics and history APIs","Protected analytics page"),
 ("ProgressSummary.jsx","Displays totals and score cards","Shows empty values safely"),
 ("PerformanceChart.jsx","Displays score trend using SVG","No additional chart library"),
 ("InterviewHistory.jsx","Provides filters, paging and report links","Previous reports accessible"),
], [3100, 3800, 2100])

heading(doc,"6.2 User Documentation")
body(doc,"This section explains how a user operates the performance analytics module after completing at least one mock interview. The Analytics page is protected, therefore the user must first register and log in.")
heading(doc,"6.2.1 Opening Performance Analytics",2)
bullets(doc,[
 "Start the backend server and frontend development server, then open the platform in the browser.",
 "Log in using a registered account.",
 "Open the Dashboard and click the Performance analytics button.",
 "The system opens the protected route: /analytics.",
 "If no interview is completed, the page shows an empty-state message and a Start new interview button.",
])
heading(doc,"6.2.2 Completing an Interview to Create Analytics",2)
bullets(doc,[
 "Generate an interview from the resume profile. Five questions are recommended for quick testing.",
 "Type an answer or select Answer using microphone, allow microphone access, speak, stop voice input and edit the transcript if needed.",
 "Select Save and continue for every question. The answer is evaluated and saved.",
 "On the final question, select Save and finish. The final performance report opens automatically.",
 "Return to the dashboard and open Performance analytics again to see the score and interview-history entry.",
])
heading(doc,"6.2.3 Understanding the Analytics Page",2)
table(doc,["Area","What the user sees","How to use it"],[
 ("Summary cards","Completed interviews, average score, highest score and improvement","Compare current progress with earlier interviews"),
 ("Score trend","One point per completed interview","Hover or read labels to compare scores over time"),
 ("Category analysis","Correctness, confidence, communication and technical accuracy","Focus practice on the weakest category"),
 ("Improvement plan","Repeated evaluation suggestions","Use frequent suggestions as a revision checklist"),
 ("Interview history","Role, type, difficulty, score, date and View action","Review any saved report"),
 ("Filters","Role search, interview type and difficulty","Narrow history results without affecting stored data"),
], [2200, 3900, 2800])
heading(doc,"6.2.4 Demonstration Procedure",2)
body(doc,"For a classroom demonstration, create and complete two short interviews with the same account. Open Performance analytics, show both score points in the graph, select a type or difficulty filter, and click View on an older interview. Refresh the browser page to demonstrate that the report and history come from MongoDB Atlas rather than temporary frontend state.")

heading(doc,"6.2.5 Screenshots of Module 6",2)
body(doc,"The following screenshots show the implemented performance analytics interface before completed interviews are available. The empty-state design clearly guides a new user to complete an interview, while preserving the same summary, graph, category and filtering layout used for completed data.")
screenshot(doc, Path(r"C:\Users\Faiz6\AppData\Local\Temp\codex-clipboard-cc952bbc-c08b-4a6d-8e80-fb2e93f26041.png"), "Figure 6.1: Performance Analytics overview - summary cards, score trend and category analysis.")
screenshot(doc, Path(r"C:\Users\Faiz6\AppData\Local\Temp\codex-clipboard-298182a9-b02d-4568-bce8-b0934136abd7.png"), "Figure 6.2: Improvement plan and Previous Interviews section with role, type and difficulty filters.")

heading(doc,"6.3 Conclusion")
body(doc,"Module 6 completes the user-progress layer of the AI Interview Preparation Platform. The results confirm that completed interviews are stored permanently, shown only to their owner and transformed into meaningful scores, trends and recommendations. The module gives users a practical way to track improvement and prepares the application for final deployment and security improvements in Module 7.")
doc.core_properties.title="Chapter 6 - Results and Discussion"; doc.core_properties.author="Student"; doc.core_properties.subject="AI Interview Preparation Platform - Module 6"
doc.save(OUTPUT)
print(OUTPUT)
